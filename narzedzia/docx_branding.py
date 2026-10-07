#!/usr/bin/env python3
"""Dodaje do ankiet DOCX nagłówek z logotypami (TUZ, COVE Polska, WIN4SMEs, UE)
oraz stopkę z klauzulą o dofinansowaniu UE.

Źródła: narzedzia/zrodla/*.docx  →  wynik: materialy/ankiety/*.docx
"""
import pathlib
import sys

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

ROOT = pathlib.Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "narzedzia"))
from branding import DISCLAIMER, LICENSE  # noqa: E402

A = ROOT / "asets" / "druk"
LOGOS = [("tuz.png", 12, WD_ALIGN_PARAGRAPH.LEFT), ("cove-polska.png", 7.5, WD_ALIGN_PARAGRAPH.CENTER),
         ("win4smes.png", 11, WD_ALIGN_PARAGRAPH.CENTER), ("eu-dofinansowanie.png", 7.5, WD_ALIGN_PARAGRAPH.RIGHT)]
WIDTHS = [30, 50, 35, 60]


def no_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tbl_pr.append(borders)


def bottom_rule(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    el = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "4"), ("space", "4"), ("color", "B9C6CC")):
        el.set(qn(f"w:{k}"), v)
    bdr.append(el)
    p_pr.append(bdr)


def brand(src, dst):
    doc = Document(src)
    for section in doc.sections:
        section.header_distance = Mm(8)
        section.footer_distance = Mm(8)
        if section.top_margin < Mm(30):
            section.top_margin = Mm(30)
        header = section.header
        for p in header.paragraphs:
            p.text = ""
        width = section.page_width - section.left_margin - section.right_margin
        table = header.add_table(1, len(LOGOS), width)
        no_borders(table)
        for cell, (fname, h, align), w in zip(table.rows[0].cells, LOGOS, WIDTHS):
            cell.width = Mm(w)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            par = cell.paragraphs[0]
            par.alignment = align
            par.add_run().add_picture(str(A / fname), height=Mm(h))
        rule = header.add_paragraph()
        bottom_rule(rule)
        rule.paragraph_format.space_after = Pt(0)
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.text = ""
        run = fp.add_run(f"{DISCLAIMER} {LICENSE}")
        run.font.size = Pt(7)
        run.font.color.rgb = RGBColor(0x46, 0x61, 0x6B)
    doc.save(dst)


if __name__ == "__main__":
    for src in sorted((ROOT / "narzedzia" / "zrodla").glob("*.docx")):
        dst = ROOT / "materialy" / "ankiety" / src.name
        brand(src, dst)
        print(dst.relative_to(ROOT))
